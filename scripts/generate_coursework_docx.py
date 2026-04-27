from __future__ import annotations

from datetime import date
from pathlib import Path


def _add_field(paragraph, field_code: str) -> None:
    """Insert a Word field code into a paragraph.

    Note: Word updates fields on open or via Update Fields.
    """

    from docx.oxml import OxmlElement  # type: ignore
    from docx.oxml.ns import qn  # type: ignore

    run = paragraph.add_run()

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code

    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    r = run._r
    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)
    r.append(fld_end)


def _set_paragraph_format(p, *, indent_cm: float = 1.25) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.shared import Cm, Pt  # type: ignore

    pf = p.paragraph_format
    pf.first_line_indent = Cm(indent_cm)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


def _add_text(doc, text: str) -> None:
    for para in text.strip().split("\n\n"):
        p = doc.add_paragraph(para.strip())
        _set_paragraph_format(p)


def _setup_styles(doc) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.shared import Cm, Pt  # type: ignore
    from docx.oxml.ns import qn  # type: ignore

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)

    # Ensure font mapping for Word (including East Asian run)
    rfonts = style._element.rPr.rFonts  # type: ignore[attr-defined]
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "Times New Roman")

    # Heading 1 (chapter titles)
    h1 = doc.styles["Heading 1"]
    h1.font.name = "Times New Roman"
    h1.font.size = Pt(14)
    h1.font.bold = True
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(12)

    # Heading 2 (subsections)
    h2 = doc.styles["Heading 2"]
    h2.font.name = "Times New Roman"
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.first_line_indent = Cm(0)


def _setup_page(doc) -> None:
    from docx.shared import Cm  # type: ignore

    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)

    # Page numbers: do not show on title page
    section.different_first_page_header_footer = True

    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_field(p, " PAGE ")


def _title_page(doc) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.shared import Pt  # type: ignore

    # First page footer stays empty
    doc.sections[0].first_page_footer.is_linked_to_previous = True

    def add_center(line: str, *, bold: bool = False, size: int = 14) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.bold = bold
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(0)

    def add_right(line: str) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(line)
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        p.paragraph_format.space_after = Pt(0)

    add_center("(УКАЖИТЕ ПОЛНОЕ НАЗВАНИЕ ВУЗА)", bold=True)
    add_center("(ИНСТИТУТ / ФАКУЛЬТЕТ)")
    add_center("(КАФЕДРА)")
    doc.add_paragraph("")
    doc.add_paragraph("")
    add_center("КУРСОВАЯ РАБОТА", bold=True)
    add_center("по теме:")
    add_center(
        "Доменно-ориентированный RAG для правовой информации: разработка и оценка качества ответов",
        bold=True,
    )
    doc.add_paragraph("")
    doc.add_paragraph("")

    add_right("Выполнил(а): студент(ка) (ФИО), группа (ГРУППА)")
    add_right("Руководитель: (ДОЛЖНОСТЬ, УЧ. СТЕПЕНЬ) (ФИО)")

    # Push city/year to bottom with blank paragraphs
    for _ in range(8):
        doc.add_paragraph("")

    add_center(f"(ГОРОД) {date.today().year}")


def _toc(doc) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore

    p = doc.add_paragraph("СОДЕРЖАНИЕ")
    p.style = doc.styles["Heading 1"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    toc_p = doc.add_paragraph()
    _add_field(toc_p, ' TOC \\o "1-3" \\h \\z \\u ')
    doc.add_page_break()


def _coursework_body(doc) -> None:
    # Введение
    doc.add_paragraph("ВВЕДЕНИЕ", style="Heading 1")
    _add_text(
        doc,
        """
Актуальность темы определяется активным внедрением больших языковых моделей (LLM) в прикладные информационные системы, включая консультирование пользователей по специализированным вопросам. Для правового домена особенно критичны проверяемость ответа, корректность формулировок и опора на действующие нормативные документы. При этом LLM, работающие без доступа к источникам, могут допускать галлюцинации — правдоподобные, но неверные утверждения, что недопустимо в юридически значимых сценариях.

Подход Retrieval-Augmented Generation (RAG) снижает риск галлюцинаций за счет генерации ответа на основе релевантных фрагментов документов, найденных в корпусе. Однако качество RAG зависит от нескольких факторов: качества разбиения документов на фрагменты (чанки), точности поиска (retrieval), отбора контекста, стратегии промптинга и надежности механизма оценки.

Целью данной курсовой работы является разработка и экспериментальная оценка доменно-ориентированной RAG-системы для правовой информации на базе реализованного проекта (Telegram-бота), выполняющего поиск по локальному корпусу нормативных документов и формирование ответа на вопрос пользователя.

Для достижения цели решаются следующие задачи:
1) рассмотреть теоретические основы доменно-ориентированного RAG и типовые риски в юридическом домене;
2) описать архитектуру разработанного решения, включая построение индекса, поиск и генерацию ответа;
3) подготовить контрольный (gold) набор вопросов и ответов и определить протокол экспериментов;
4) провести серию улучшений (гипотез) в части retrieval и prompt-инжиниринга;
5) выполнить оценку качества по метрикам RAGAS и выбрать конфигурацию, максимизирующую метрику Faithfulness.

Объект исследования — система вопросно-ответного взаимодействия по доменно-ограниченному корпусу правовых документов.

Предмет исследования — методы повышения фактической обоснованности ответов в RAG-системе (retrieval, reranking, промптинг) и их количественная оценка.

Методы исследования: анализ литературы по RAG, разработка программного прототипа, экспериментальное сравнение конфигураций, оценка результатов метриками RAGAS.

Практическая значимость работы заключается в получении воспроизводимой реализации юридического RAG-бота и набора экспериментальных рекомендаций по повышению Faithfulness (обоснованности) ответа на основе внедрения cross-encoder reranker и «faithfulness-first» промпта.
""",
    )
    doc.add_page_break()

    # Глава 1
    doc.add_paragraph(
        "1 Теоретические основы доменно-ориентированного RAG", style="Heading 1"
    )
    doc.add_paragraph("1.1 Ограничения LLM в юридическом домене", style="Heading 2")
    _add_text(
        doc,
        """
Юридический домен характеризуется высокой ценой ошибки: некорректная трактовка нормы или ссылка на несуществующий пункт может привести к неверным решениям пользователя. Большие языковые модели обучаются на больших массивов текстов и оптимизируются на правдоподобность ответа, но не гарантируют истинность фактов. В результате возникают галлюцинации: модель может «достраивать» недостающие детали, смешивать разные источники или переносить знания из других контекстов.

В условиях доменно-ориентированного ассистента ключевыми требованиями являются: (а) минимизация недоказуемых утверждений, (б) воспроизводимость ответа относительно исходных документов, (в) возможность предоставления подтверждения (цитаты/ссылки на фрагменты), (г) контроль отказа: если сведений нет, система должна корректно сообщать об этом.
""",
    )

    doc.add_paragraph("1.2 Архитектура RAG: retrieval и generation", style="Heading 2")
    _add_text(
        doc,
        """
RAG-система в общем случае включает этапы: подготовка корпуса (ingestion), индексирование, извлечение релевантного контекста (retrieval) и генерация ответа (generation). На этапе ingestion документы нормализуются, разбиваются на фрагменты и снабжаются метаданными (например, имя файла, идентификатор документа, номера пунктов). На этапе индексирования формируются векторные представления фрагментов и строится индекс для быстрого поиска.

Retrieval может быть реализован как семантический поиск по эмбеддингам (dense retrieval) и/или как лексический поиск (sparse retrieval, например BM25). Практика показывает, что гибридный подход позволяет компенсировать слабые стороны каждого метода: семантика лучше для перефразированных запросов, BM25 — для точных терминов, аббревиатур и ссылок.

Generation — это формирование ответа LLM на основе вопроса и извлеченного контекста. В доменно-ориентированном RAG важен строгий промпт, ограничивающий модель рамками контекста, а также правила отказа при отсутствии релевантных фрагментов.
""",
    )

    doc.add_paragraph("1.3 Метрики качества RAGAS", style="Heading 2")
    _add_text(
        doc,
        """
Для количественной оценки качества используется библиотека RAGAS, предоставляющая метрики, ориентированные на сценарии RAG. В работе применяются три метрики:

Answer Relevance — оценивает, насколько сгенерированный ответ соответствует пользовательскому вопросу.

Faithfulness (Groundedness) — оценивает, насколько утверждения в ответе подтверждаются извлеченным контекстом. Для юридического домена эта метрика является приоритетной, поскольку снижает риск недоказуемых и потенциально ложных утверждений.

Context Relevance — оценивает, насколько извлеченный контекст релевантен вопросу, то есть насколько retrieval обеспечивает полезные для ответа фрагменты.

Оценивание проводится при помощи LLM-судьи и эмбеддингов для семантических сопоставлений, при этом важно обеспечить воспроизводимость условий: одинаковый корпус, одинаковый протокол формирования контекста и фиксированные параметры генерации.
""",
    )

    doc.add_paragraph(
        "1.4 Инструменты: FAISS, BM25, reranker, Ollama", style="Heading 2"
    )
    _add_text(
        doc,
        """
FAISS — библиотека для эффективного поиска ближайших соседей во векторных пространствах; в проекте используется индекс по косинусной близости (через скалярное произведение нормированных эмбеддингов).

BM25 — классический метод информационного поиска на основе статистики терминов, удобен как дополняющий компонент к семантическому поиску.

Reranker (cross-encoder) — модель, которая оценивает релевантность пары (вопрос, фрагмент) более точно, чем простая близость эмбеддингов, поскольку обрабатывает пару совместно. Reranker используется для повторного ранжирования ограниченного набора кандидатов, полученных на этапе retrieval.

Ollama — локальный сервер для запуска LLM; в проекте он используется для генерации ответов и для LLM-судьи при оценке RAGAS без внешних API.
""",
    )
    doc.add_page_break()

    # Глава 2
    doc.add_paragraph("2 Проектирование и реализация системы", style="Heading 1")
    doc.add_paragraph(
        "2.1 Постановка задачи юридического ассистента", style="Heading 2"
    )
    _add_text(
        doc,
        """
Разрабатываемая система должна отвечать на вопросы пользователя по локальному корпусу нормативных и регламентирующих документов. Выход системы — краткий, прямой ответ и список источников (файлы и идентификаторы пунктов), что позволяет пользователю проверить ответ.

Ключевое ограничение: система не должна добавлять информацию, отсутствующую в контексте. Если релевантных сведений нет, система обязана сообщить об отсутствии данных в предоставленном контексте.
""",
    )

    doc.add_paragraph("2.2 Подготовка корпуса документов и чанкинг", style="Heading 2")
    _add_text(
        doc,
        """
Корпус документов хранится локально и проходит этап разбиения на фрагменты (чанки). Для каждого чанка сохраняются: текст, уникальный идентификатор, а также метаданные (например, исходный файл, идентификатор документа, номера пунктов и разделов). Наличие метаданных важно для дальнейшего формирования источников и для контроля качества retrieval.

При разбиении необходимо балансировать между полнотой и точностью: слишком длинные чанки ухудшают точность поиска и увеличивают контекст, слишком короткие могут терять смысловые связи и приводить к неполным ответам.
""",
    )

    doc.add_paragraph(
        "2.3 Индексация и retrieval (FAISS, hybrid BM25)", style="Heading 2"
    )
    _add_text(
        doc,
        """
В проекте реализован семантический поиск по эмбеддингам. Для векторизации используется модель семейства E5 (multilingual), обеспечивающая поддержку русского языка. Векторный индекс хранится в FAISS.

В качестве улучшения рассматривается гибридный retrieval: к семантическим кандидатам добавляется оценка BM25, и далее формируется итоговый скор как взвешенная комбинация. Это повышает устойчивость к запросам, содержащим редкие термины и устойчивые формулировки.

Для максимизации Faithfulness в данной версии также внедрен этап reranking: после получения пула кандидатов (например, top-50) применяется cross-encoder, переоценивающий релевантность каждого кандидата к вопросу. На выход подается top-k переранжированных фрагментов.
""",
    )

    doc.add_paragraph(
        "2.4 Генерация ответа и управление источниками", style="Heading 2"
    )
    _add_text(
        doc,
        """
Генерация ответа выполняется локальной LLM через Ollama. Для снижения галлюцинаций применяется «faithfulness-first» промпт: модель должна использовать только факты, прямо содержащиеся в контексте, и добавить дословные короткие цитаты в качестве подтверждения. Если в контексте нет прямой информации, модель возвращает строго заданную фразу отказа.

Отдельно решается проблема источников: вместо просьбы к LLM «сгенерировать источники» (что может приводить к выдуманным ссылкам) источники формируются программно по метаданным выбранных чанков и добавляются к ответу на стороне приложения.
""",
    )

    doc.add_paragraph("2.5 Интеграция в Telegram-бота", style="Heading 2")
    _add_text(
        doc,
        """
Пользовательский интерфейс реализован в виде Telegram-бота. После получения сообщения бот запускает цепочку: retrieval → сборка контекста → генерация ответа → добавление источников и отправка результата пользователю. Для управления параметрами (top-k, включение reranker) используются переменные окружения.
""",
    )
    doc.add_page_break()

    # Глава 3
    doc.add_paragraph("3 Экспериментальная оценка и результаты", style="Heading 1")
    doc.add_paragraph("3.1 Датасет и протокол эксперимента", style="Heading 2")
    _add_text(
        doc,
        """
Для оценки качества сформирован gold-датасет на основе таблицы вопросов. Для корректности эксперимента оставлены только вопросы, для которых исходные документы присутствуют в локальном корпусе (in-scope). Итоговый набор содержит 70 вопросов.

Оценивание выполняется с помощью RAGAS по трём метрикам: Answer Relevance, Faithfulness, Context Relevance. В качестве LLM-судьи используется локальная модель Qwen через Ollama. Для воспроизводимости температура генерации фиксируется на 0.0.
""",
    )

    doc.add_paragraph("3.2 Baseline и гипотезы улучшения", style="Heading 2")
    _add_text(
        doc,
        """
В работе проверялись гипотезы улучшения качества:
1) программное формирование источников (LLM не генерирует список ссылок);
2) адаптивный выбор top-k для retrieval;
3) гибридный retrieval FAISS + BM25;
4) внедрение cross-encoder reranker для переранжирования кандидатов;
5) «faithfulness-first» промпт для уменьшения домыслов.

Основной критерий выбора итоговой конфигурации — максимизация Faithfulness при приемлемом уровне релевантности ответа.
""",
    )

    doc.add_paragraph("3.3 Результаты оценки RAGAS", style="Heading 2")

    from docx.shared import Pt  # type: ignore

    p = doc.add_paragraph(
        "В таблице 1 приведены результаты ключевых прогонов (средние значения метрик)."
    )
    _set_paragraph_format(p)

    # Table: results
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Конфигурация"
    hdr[1].text = "Answer Relevance"
    hdr[2].text = "Faithfulness"
    hdr[3].text = "Context Relevance"

    rows = [
        (
            "Baseline (ручные ответы бота, in-scope)",
            "0.7485",
            "0.5341",
            "0.8250",
        ),
        (
            "FAISS + rerank (без faithfulness-first промпта)",
            "0.7238",
            "0.5827",
            "0.8786",
        ),
        (
            "FAISS + rerank + faithfulness-first промпт (итог)",
            "0.7361",
            "0.7705",
            "0.8750",
        ),
    ]
    for r in rows:
        c = table.add_row().cells
        c[0].text, c[1].text, c[2].text, c[3].text = r

    cap = doc.add_paragraph("Таблица 1 — Сравнение конфигураций по метрикам RAGAS")
    cap.paragraph_format.first_line_indent = None
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(12)

    doc.add_paragraph("3.4 Анализ результатов", style="Heading 2")
    _add_text(
        doc,
        """
Эксперименты показывают, что наибольший прирост Faithfulness достигается при сочетании двух факторов: улучшения качества контекста (за счет reranking) и ограничений на генерацию (faithfulness-first промпт). Reranker повышает вероятность того, что в контекст попадут фрагменты, непосредственно отвечающие на вопрос, а строгий промпт снижает склонность модели «обобщать» и делать недоказуемые выводы.

Введение дословных цитат выполняет двойную роль: (а) дисциплинирует модель, вынуждая опираться на текст, (б) повышает доверие пользователя, поскольку подтверждение видно непосредственно в ответе.

При этом следует учитывать компромисс: более строгий промпт может снизить Answer Relevance в случаях, когда корректный ответ требует синтеза информации из нескольких мест или аккуратной интерпретации. В рамках юридического ассистента выбран приоритет на Faithfulness как на ключевую метрику безопасности.
""",
    )

    doc.add_paragraph("3.5 Ограничения и риски", style="Heading 2")
    _add_text(
        doc,
        """
Ограничения решения:
1) Зависимость от качества разметки и метаданных (если номера пунктов определены неверно, источники могут быть неполными).
2) Ограниченный контекст: при больших документах релевантные фрагменты могут не поместиться в окно контекста.
3) Стоимость reranker: cross-encoder требует больше вычислений, чем простое сравнение эмбеддингов, что увеличивает задержку ответа.
4) Оценка RAGAS зависит от LLM-судьи и может содержать шум; необходимы повторные прогоны и анализ отдельных примеров.

Для дальнейшего развития возможно: улучшение разбиения документов, обучение/подбор доменно-специфичного reranker, внедрение кэширования, расширение gold-датасета и проведение пользовательского тестирования.
""",
    )
    doc.add_page_break()

    # Заключение
    doc.add_paragraph("ЗАКЛЮЧЕНИЕ", style="Heading 1")
    _add_text(
        doc,
        """
В курсовой работе разработана доменно-ориентированная RAG-система для правовой информации на базе локального корпуса документов и Telegram-интерфейса. Рассмотрены теоретические основы RAG, ограничения LLM в юридическом домене и подходы к оценке качества.

В практической части описана архитектура проекта: подготовка документов и чанков, индексирование во FAISS, retrieval, генерация ответа локальной LLM через Ollama и программное формирование источников. Проведена экспериментальная оценка по метрикам RAGAS на in-scope gold-датасете из 70 вопросов.

Основной результат — подтверждение гипотезы, что сочетание cross-encoder reranker и «faithfulness-first» промпта существенно повышает Faithfulness (до 0.7705 в среднем), что является критически важным показателем для юридического ассистента. Полученное решение может быть использовано как основа для дальнейших улучшений качества и надежности доменно-ориентированных RAG-систем.
""",
    )
    doc.add_page_break()

    # Список литературы
    doc.add_paragraph("СПИСОК ЛИТЕРАТУРЫ", style="Heading 1")
    refs = [
        "1. Lewis P., Perez E., Piktus A. et al. Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. 2020.",
        "2. RAGAS: Automated Evaluation of Retrieval Augmented Generation. Documentation and examples. 2024.",
        "3. Johnson J., Douze M., Jégou H. Billion-scale similarity search with GPUs (FAISS). 2017.",
        "4. Robertson S., Zaragoza H. The Probabilistic Relevance Framework: BM25 and Beyond. 2009.",
        "5. Reimers N., Gurevych I. Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks. 2019.",
        "6. Sentence-Transformers: Cross-Encoder documentation (reranking). 2023–2025.",
        "7. Qwen2.5: Technical report and model card. 2024.",
        "8. Ollama: Local LLM server. Documentation. 2024–2026.",
        "9. Telegram Bot API. Documentation. 2024–2026.",
        "10. ГОСТ 7.32–2017. Отчет о научно-исследовательской работе. Структура и правила оформления.",
        "11. ГОСТ Р 7.0.5–2008. Библиографическая ссылка. Общие требования и правила составления.",
    ]
    for r in refs:
        p = doc.add_paragraph(r)
        _set_paragraph_format(p, indent_cm=0.0)

    doc.add_page_break()

    # Приложения
    doc.add_paragraph("ПРИЛОЖЕНИЕ А (справочное)", style="Heading 1")
    doc.add_paragraph("Команды запуска оценки", style="Heading 2")
    _add_text(
        doc,
        """
Пример запуска оценки RAGAS с генерацией ответов и включенным reranker:

uv run python scripts/eval_ragas_stage1.py --input-xlsx data/gold/gold_in_scope.xlsx \
  --out-xlsx data/gold/gen_faiss_rerank_faithfulprompt_scored_full.xlsx \
  --index-dir data/faiss_index --index-name moscow_kb --generate-answers \
  --use-reranker --reranker-candidates-k 50
""",
    )

    doc.add_paragraph("ПРИЛОЖЕНИЕ Б (справочное)", style="Heading 1")
    doc.add_paragraph("Переменные окружения бота", style="Heading 2")
    _add_text(
        doc,
        """
Основные параметры конфигурации:

TELEGRAM_BOT_TOKEN — токен Telegram-бота.
TOP_K — число фрагментов контекста.
USE_RERANKER — включение cross-encoder reranker (1/0).
RERANKER_MODEL — имя модели reranker.
RERANKER_CANDIDATES_K — размер пула кандидатов для переранжирования.
OLLAMA_BASE_URL, OLLAMA_MODEL, OLLAMA_TIMEOUT — параметры доступа к Ollama.
""",
    )


def main() -> int:
    from docx import Document  # type: ignore

    out_path = Path("cursovay.docx").resolve()

    doc = Document()
    _setup_styles(doc)
    _setup_page(doc)
    _title_page(doc)
    doc.add_page_break()
    _toc(doc)
    _coursework_body(doc)

    doc.save(str(out_path))
    print(f"Saved: {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
