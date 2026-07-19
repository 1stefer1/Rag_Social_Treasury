# uv run python scripts/eval_ragas_stage1.py
from __future__ import annotations

import argparse
import asyncio
import json
import math
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from statistics import mean
from typing import Any, Dict, List, Optional

sys.path.append(str(Path(__file__).resolve().parents[1]))

from datasets import Dataset
from elasticsearch import Elasticsearch
from openpyxl import Workbook, load_workbook
from qdrant_client import QdrantClient
from ragas import evaluate
from ragas.metrics._answer_relevance import AnswerRelevancy
from ragas.metrics._context_precision import ContextPrecision
from ragas.metrics._context_recall import ContextRecall
from ragas.metrics._faithfulness import Faithfulness
from ragas.metrics._nv_metrics import ContextRelevance
from ragas.run_config import RunConfig

from src.utils.embedder import Embedder
from src.utils.eval_tracker import EvalTracker
from src.utils.generator import LLM
from src.utils.rag_pipeline import VanillaRAG
from src.utils.ragas_support import (
    E5RagasEmbeddings,
    OllamaRagasLLM,
    warmup_embedder,
)
from src.utils.retriever import Retriever
from src.utils.sparse_store import ElasticsearchSparseStore
from src.utils.vector_store import QdrantVectorStore


@dataclass
class Row:
    id: str
    question: str
    gold_answer: str
    rag_answer: str
    gold_source: str
    matched_doc: str


def _safe_str(v: Any) -> str:
    if v is None:
        return ""
    return str(v)


def _load_rows(xlsx_path: Path) -> List[Row]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    col = {str(name).strip(): i for i, name in enumerate(header) if name is not None}
    required = [
        "id",
        "question",
        "gold_answer",
        "rag_answer",
        "gold_source",
        "matched_doc",
    ]
    for r in required:
        if r not in col:
            raise KeyError(f"Missing column '{r}' in {xlsx_path}")

    rows: List[Row] = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        row = Row(
            id=_safe_str(r[col["id"]]).strip(),
            question=_safe_str(r[col["question"]]).strip(),
            gold_answer=_safe_str(r[col["gold_answer"]]).strip(),
            rag_answer=_safe_str(r[col["rag_answer"]]).strip(),
            gold_source=_safe_str(r[col["gold_source"]]).strip(),
            matched_doc=_safe_str(r[col["matched_doc"]]).strip(),
        )
        if not row.question:
            continue
        rows.append(row)
    return rows


def _to_float(v: Any) -> Optional[float]:
    if v is None:
        return None
    try:
        x = float(v)
        if math.isnan(x) or math.isinf(x):
            return None
        return x
    except Exception:
        return None


def _mean(values: List[Optional[float]]) -> Optional[float]:
    xs = [v for v in values if v is not None]
    if not xs:
        return None
    return mean(xs)


def _percentile(values: List[Optional[float]], q: float) -> Optional[float]:
    xs = sorted(v for v in values if v is not None)
    if not xs:
        return None
    if len(xs) == 1:
        return xs[0]
    pos = (len(xs) - 1) * q
    lo = math.floor(pos)
    hi = math.ceil(pos)
    if lo == hi:
        return xs[lo]
    weight = pos - lo
    return xs[lo] + (xs[hi] - xs[lo]) * weight


def _normalize_metric_name(name: str) -> str:
    mapping = {
        "answer_relevancy": "answer_relevance",
        "nv_context_relevance": "context_relevance",
    }
    return mapping.get(name, name)


def _infer_retriever_name(args: argparse.Namespace) -> str:
    if args.retriever_name:
        return args.retriever_name
    parts = ["qdrant", "elasticsearch", "rrf"]
    if args.use_reranker:
        parts.append("rerank")
    return "+".join(parts)


def _build_run_name(args: argparse.Namespace) -> str:
    if args.hypothesis_name:
        return args.hypothesis_name
    return _infer_retriever_name(args)


def _build_eval_records(
    input_rows: List[Row],
    answers: List[str],
    contexts: List[List[str]],
    scores: List[Dict[str, Any]],
    telemetry: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []
    for row, ans, ctxs, sc, tm in zip(input_rows, answers, contexts, scores, telemetry):
        ctx_join = "\n\n---\n\n".join(ctxs)
        record: Dict[str, Any] = {
            "id": row.id,
            "question": row.question,
            "gold_answer": row.gold_answer,
            "rag_answer": ans,
            "gold_source": row.gold_source,
            "matched_doc": row.matched_doc,
            "used_top_k": len(ctxs),
            "context_chars": len(ctx_join),
            "retrieved_contexts": ctx_join,
            **tm,
        }
        for key, value in sc.items():
            record[_normalize_metric_name(key)] = value
        records.append(record)
    return records


def _summarize_records(records: List[Dict[str, Any]]) -> Dict[str, float]:
    summary: Dict[str, float] = {}
    if not records:
        return summary

    numeric_fields = [
        "faithfulness",
        "answer_relevance",
        "context_precision",
        "context_recall",
        "context_relevance",
        "retrieval_latency_ms",
        "generation_latency_ms",
        "end_to_end_latency_ms",
        "used_top_k",
        "context_chars",
    ]
    for field in numeric_fields:
        values = [_to_float(record.get(field)) for record in records]
        avg = _mean(values)
        if avg is not None:
            suffix = "mean"
            if field.endswith("_ms"):
                summary[f"{field[:-3]}_mean_ms"] = avg
                p95 = _percentile(values, 0.95)
                if p95 is not None:
                    summary[f"{field[:-3]}_p95_ms"] = p95
            else:
                if field in {
                    "faithfulness",
                    "answer_relevance",
                    "context_precision",
                    "context_recall",
                    "context_relevance",
                }:
                    summary[field] = avg
                summary[f"{field}_{suffix}"] = avg

    summary["rows_evaluated"] = float(len(records))
    return summary


def _write_output(
    out_path: Path,
    records: List[Dict[str, Any]],
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "scored"

    header = list(records[0].keys()) if records else []
    ws.append(header)

    for record in records:
        ws.append([record.get(col) for col in header])

    # summary sheet
    ws2 = wb.create_sheet("summary")
    ws2.append(["metric", "mean"])
    summary = _summarize_records(records)
    for key, value in summary.items():
        ws2.append([key, value])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


async def main() -> int:
    ap = argparse.ArgumentParser(
        description="Stage-1 RAGAS evaluation using local Ollama Qwen as judge (no external API)."
    )
    ap.add_argument(
        "--input-xlsx",
        default="data/gold/gold_in_scope.xlsx",
        help="Input gold xlsx with filled rag_answer (default: %(default)s)",
    )
    ap.add_argument(
        "--generate-answers",
        action="store_true",
        help="Generate rag answers via current RAG pipeline (ignores rag_answer column)",
    )
    ap.add_argument(
        "--out-xlsx",
        default="data/gold/gold_in_scope_scored_qwen.xlsx",
        help="Output xlsx with metrics (default: %(default)s)",
    )
    ap.add_argument(
        "--qdrant-url",
        default="http://localhost:6333",
        help="Qdrant URL (default: %(default)s)",
    )
    ap.add_argument(
        "--qdrant-collection",
        default="moscow_kb",
        help="Qdrant collection (default: %(default)s)",
    )
    ap.add_argument("--top-k", type=int, default=5)
    ap.add_argument("--max-context-chars", type=int, default=14000)
    ap.add_argument("--es-url", default="http://localhost:9200")
    ap.add_argument("--es-index", default="kb_chunks")
    ap.add_argument("--dense-candidates-k", type=int, default=50)
    ap.add_argument("--sparse-candidates-k", type=int, default=50)
    ap.add_argument("--rrf-k", type=int, default=60)
    ap.add_argument("--use-reranker", action="store_true", help="Enable cross-encoder reranking")
    ap.add_argument(
        "--reranker-model",
        default="cross-encoder/mmarco-mMiniLMv2-L12-H384-v1",
        help="Reranker model name (default: %(default)s)",
    )
    ap.add_argument(
        "--reranker-candidates-k",
        type=int,
        default=50,
        help="How many top candidates to rerank (default: %(default)s)",
    )
    ap.add_argument("--limit", type=int, default=0, help="Evaluate only first N rows")
    ap.add_argument("--ollama-url", default="http://localhost:11434")
    ap.add_argument("--ollama-model", default="qwen2.5:7b-instruct")
    ap.add_argument("--tracking-uri", default="sqlite:///mlflow.db")
    ap.add_argument("--experiment-name", default="rag-evals")
    ap.add_argument("--hypothesis-name", default="")
    ap.add_argument("--prompt-version", default="vanilla_rag_v1")
    ap.add_argument("--chunk-size", type=int, default=2500)
    ap.add_argument("--embedding-model", default="intfloat/multilingual-e5-base")
    ap.add_argument("--llm-model", default="qwen2.5:7b-instruct")
    ap.add_argument("--retriever-name", default="")
    ap.add_argument("--history-json", default="")
    args = ap.parse_args()

    in_path = Path(args.input_xlsx)
    out_path = Path(args.out_xlsx)

    rows = _load_rows(in_path)
    if not args.generate_answers:
        rows = [r for r in rows if r.rag_answer]
    if args.limit and args.limit > 0:
        rows = rows[: args.limit]

    if not rows:
        raise RuntimeError("No rows to evaluate (check rag_answer column)")

    tracker = EvalTracker(
        tracking_uri=args.tracking_uri,
        experiment_name=args.experiment_name,
        run_name=_build_run_name(args),
        tags={
            "hypothesis_name": args.hypothesis_name or _build_run_name(args),
            "prompt_version": args.prompt_version,
            "run_origin": "eval_script",
        },
    )
    tracker.start()
    try:
        embedder = Embedder(model_name=args.embedding_model)
        await warmup_embedder(embedder)

        retriever = Retriever(
            embedder,
            QdrantVectorStore(QdrantClient(url=args.qdrant_url), args.qdrant_collection),
            ElasticsearchSparseStore(Elasticsearch(args.es_url, request_timeout=30), args.es_index),
            top_k=args.top_k,
            dense_candidates_k=int(args.dense_candidates_k),
            sparse_candidates_k=int(args.sparse_candidates_k),
            rrf_k=int(args.rrf_k),
            use_reranker=bool(args.use_reranker),
            reranker_model=str(args.reranker_model),
            reranker_candidates_k=int(args.reranker_candidates_k),
        )
        rag: VanillaRAG | None = None
        if args.generate_answers:
            llm = LLM(model=args.llm_model)
            rag = VanillaRAG(
                retriever,
                llm,
                default_top_k=args.top_k,
                max_context_chars=args.max_context_chars,
            )

        # Build contexts + (optional) generate answers
        contexts: List[List[str]] = []
        answers: List[str] = []
        telemetry: List[Dict[str, Any]] = []

        params = {
            "input_xlsx": str(in_path),
            "qdrant_url": args.qdrant_url,
            "qdrant_collection": args.qdrant_collection,
            "es_url": args.es_url,
            "es_index": args.es_index,
            "chunk_size": args.chunk_size,
            "top_k": args.top_k,
            "max_context_chars": args.max_context_chars,
            "retriever": _infer_retriever_name(args),
            "embedding_model": embedder.model_name,
            "prompt_version": args.prompt_version,
            "llm_model": args.llm_model,
            "judge_model": args.ollama_model,
            "hypothesis_name": args.hypothesis_name or _build_run_name(args),
            "dense_candidates_k": args.dense_candidates_k,
            "sparse_candidates_k": args.sparse_candidates_k,
            "rrf_k": args.rrf_k,
            "use_reranker": args.use_reranker,
            "reranker_model": args.reranker_model,
            "reranker_candidates_k": args.reranker_candidates_k,
            "generate_answers": args.generate_answers,
            "rows_limit": args.limit,
        }
        tracker.log_params(params)

        for r in rows:
            row_started_at = time.perf_counter()
            # Retrieve once to keep contexts and generation aligned
            retrieval_started_at = time.perf_counter()
            chunks = await retriever.aretrieve(r.question, top_k=args.top_k)
            retrieval_latency_ms = (time.perf_counter() - retrieval_started_at) * 1000.0
            # Build contexts (same locator style as pipeline)
            ctxs: List[str] = []
            total = 0
            for i, c in enumerate(chunks, 1):
                from src.utils.ragas_support import format_locator

                loc = format_locator(i, c.meta or {})
                block = f"{loc}\n{c.text}".strip()
                if total + len(block) > args.max_context_chars:
                    break
                ctxs.append(block)
                total += len(block)
            contexts.append(ctxs)

            generation_latency_ms: Optional[float] = None
            if args.generate_answers:
                assert rag is not None
                prompt, used_chunks = rag._build_prompt(r.question, chunks)
                generation_started_at = time.perf_counter()
                ans = await rag.llm.arun(prompt, temperature=0.0, max_tokens=700)
                generation_latency_ms = (time.perf_counter() - generation_started_at) * 1000.0
                answers.append((ans or "").strip())
            else:
                answers.append(r.rag_answer)

            telemetry.append(
                {
                    "retrieval_latency_ms": retrieval_latency_ms,
                    "generation_latency_ms": generation_latency_ms,
                    "end_to_end_latency_ms": (time.perf_counter() - row_started_at) * 1000.0,
                }
            )

        dataset = Dataset.from_dict(
            {
                "user_input": [r.question for r in rows],
                "response": answers,
                "reference": [r.gold_answer for r in rows],
                "retrieved_contexts": contexts,
            }
        )

        judge_llm = OllamaRagasLLM(model=args.ollama_model, base_url=args.ollama_url)
        judge_emb = E5RagasEmbeddings(embedder)

        run_config = RunConfig(timeout=600, max_workers=4)

        # Use fresh metric instances (avoid global singletons being mutated between runs)
        metrics = [
            AnswerRelevancy(),
            Faithfulness(),
            ContextPrecision(),
            ContextRecall(),
            ContextRelevance(),
        ]

        result = evaluate(
            dataset,
            metrics=metrics,
            llm=judge_llm,
            embeddings=judge_emb,
            run_config=run_config,
            raise_exceptions=False,
            show_progress=True,
            batch_size=4,
        )

        scores = result.scores
        records = _build_eval_records(rows, answers, contexts, scores, telemetry)
        _write_output(out_path, records)
        summary = _summarize_records(records)

        history_runs: List[Dict[str, Any]] = []
        if args.history_json:
            with Path(args.history_json).open("r", encoding="utf-8") as f:
                history_runs = json.load(f)

        # Print quick summary
        print(f"Rows evaluated: {len(rows)}")
        for key in (
            "faithfulness",
            "answer_relevance",
            "context_precision",
            "context_recall",
            "context_relevance",
        ):
            value = summary.get(key)
            if value is not None:
                print(f"{key}: mean={value:.4f}")
        print(f"Saved: {out_path}")

        tracker.log_metrics(summary)
        tracker.log_eval_payload(
            rows=records,
            params=params,
            summary_metrics=summary,
            out_xlsx=out_path,
            history_runs=history_runs,
        )
        tracker.end(status="FINISHED")
        return 0
    except Exception:
        tracker.end(status="FAILED")
        raise


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
