from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

sys.path.append(str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from src.utils.eval_tracker import EvalTracker, create_history_report


def _safe_float(value: Any) -> Optional[float]:
    try:
        if value is None:
            return None
        return float(value)
    except (TypeError, ValueError):
        return None


def _normalize_metric_name(name: str) -> str:
    mapping = {
        "answer_relevancy": "answer_relevance",
        "nv_context_relevance": "context_relevance",
    }
    return mapping.get(name, name)


def _infer_params(stem: str) -> Dict[str, Any]:
    name = stem.lower()
    retriever = "faiss"
    if "hybrid" in name or "bm25" in name:
        retriever = "faiss+bm25"
    if "rerank" in name:
        retriever += "+rerank"

    prompt_version = "vanilla_rag_v1"
    if "strict" in name:
        prompt_version = "strict_prompt"
    elif "faithfulprompt" in name:
        prompt_version = "faithfulness_first"
    elif "nosources" in name:
        prompt_version = "no_sources"

    llm_model = "qwen2.5:7b-instruct" if "baseline_manual" not in name else "manual_answers"

    top_k = 5
    if "topk" in name:
        digits = "".join(ch for ch in name.split("topk", 1)[1] if ch.isdigit())
        if digits:
            top_k = int(digits)

    return {
        "hypothesis_name": stem,
        "retriever": retriever,
        "prompt_version": prompt_version,
        "embedding_model": "intfloat/multilingual-e5-base",
        "llm_model": llm_model,
        "chunk_size": 2500,
        "top_k": top_k,
        "source": "excel_backfill",
    }


def _load_summary(path: Path) -> Optional[Dict[str, Any]]:
    if path.name.startswith("~$"):
        return None

    wb = load_workbook(path, read_only=True, data_only=True)
    if "summary" not in wb.sheetnames:
        return None

    summary_sheet = wb["summary"]
    header = [
        cell for cell in next(summary_sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    ]
    header_map = {str(name).strip(): idx for idx, name in enumerate(header) if name is not None}
    metric_idx = header_map.get("metric", 0)
    mean_idx = header_map.get("mean", 1)
    count_idx = header_map.get("count")

    metrics: Dict[str, float] = {}
    count: Optional[float] = None
    for row in summary_sheet.iter_rows(min_row=2, values_only=True):
        raw_name = row[metric_idx]
        if raw_name is None:
            continue
        metric_name = _normalize_metric_name(str(raw_name).strip())
        metric_value = _safe_float(row[mean_idx])
        if metric_value is not None:
            metrics[metric_name] = metric_value
        if count is None and count_idx is not None:
            count = _safe_float(row[count_idx])

    if count is None and "scored" in wb.sheetnames:
        scored_sheet = wb["scored"]
        count = max(scored_sheet.max_row - 1, 0)

    payload = {
        "source_file": path.name,
        "run_name": path.stem,
        "rows_evaluated": count or 0,
        **metrics,
        **_infer_params(path.stem),
    }
    return payload


def main() -> int:
    ap = argparse.ArgumentParser(
        description="Backfill historical Excel eval summaries into local MLflow and build comparison charts."
    )
    ap.add_argument("--input-dir", default="data/gold")
    ap.add_argument("--tracking-uri", default="sqlite:///mlflow.db")
    ap.add_argument("--experiment-name", default="rag-evals")
    ap.add_argument("--dashboard-run-name", default="historical-eval-dashboard")
    args = ap.parse_args()

    input_dir = Path(args.input_dir)
    paths = sorted(input_dir.glob("*.xlsx"))
    if not paths:
        raise FileNotFoundError(f"No xlsx files found in {input_dir}")

    summaries: List[Dict[str, Any]] = []
    for path in paths:
        summary = _load_summary(path)
        if summary is None:
            continue
        summaries.append(summary)

        tracker = EvalTracker(
            tracking_uri=args.tracking_uri,
            experiment_name=args.experiment_name,
            run_name=summary["run_name"],
            tags={
                "run_origin": "excel_backfill",
                "hypothesis_name": summary["hypothesis_name"],
                "source_file": summary["source_file"],
            },
        )
        tracker.start()
        tracker.log_params(
            {
                "hypothesis_name": summary["hypothesis_name"],
                "retriever": summary["retriever"],
                "embedding_model": summary["embedding_model"],
                "prompt_version": summary["prompt_version"],
                "llm_model": summary["llm_model"],
                "chunk_size": summary["chunk_size"],
                "top_k": summary["top_k"],
                "source": summary["source"],
            }
        )
        tracker.log_metrics(
            {
                key: value
                for key, value in summary.items()
                if key
                in {
                    "faithfulness",
                    "answer_relevance",
                    "context_precision",
                    "context_recall",
                    "context_relevance",
                    "rows_evaluated",
                }
            }
        )
        tracker.log_artifact_file(path, artifact_path="historical_xlsx")
        tracker.end(status="FINISHED")

    if not summaries:
        raise RuntimeError("No historical summaries with a 'summary' sheet were found.")

    dashboard_tracker = EvalTracker(
        tracking_uri=args.tracking_uri,
        experiment_name=args.experiment_name,
        run_name=args.dashboard_run_name,
        tags={
            "run_origin": "historical_dashboard",
        },
    )
    dashboard_tracker.start()
    dashboard_tracker.log_params({"history_run_count": len(summaries)})
    dashboard_tracker.log_metrics({"rows_evaluated": float(len(summaries))})

    import tempfile

    with tempfile.TemporaryDirectory(prefix="mlflow-history-") as tmp_dir_str:
        report_dir = Path(tmp_dir_str) / "history"
        create_history_report(summaries, report_dir)
        dashboard_tracker.log_artifact_dir(report_dir, artifact_path="history")

    dashboard_tracker.end(status="FINISHED")

    print(f"Imported historical runs: {len(summaries)}")
    print(f"Tracking URI: {args.tracking_uri}")
    print(f"Experiment: {args.experiment_name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
