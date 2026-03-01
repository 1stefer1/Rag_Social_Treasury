from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Optional

from openpyxl import Workbook, load_workbook


@dataclass(frozen=True)
class DocKey:
    doc_type: str
    date: str  # dd.mm.yyyy
    number: str


_DATE_RE = re.compile(r"\b(\d{2}\.\d{2}\.\d{4})\b")
_NUM_RE = re.compile(
    r"(?i)(?:№|\bN\b)\s*([0-9]+(?:[-–—][0-9]+)?(?:[-–—][A-Za-zА-Яа-я0-9]+)?)"
)
_OT_DATE_RE = re.compile(r"(?i)\bот\s*(\d{2}\.\d{2}\.\d{4})\b")


def _norm_spaces(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "")).strip()


def _norm_dashes(s: str) -> str:
    return (s or "").replace("–", "-").replace("—", "-")


def _detect_doc_type(text: str) -> Optional[str]:
    t = _norm_spaces(text).lower()
    t = t.replace("№", "n")

    if "постановлен" in t and "правительств" in t:
        if "москв" in t:
            return "pp_moscow"
        if "российск" in t or "рф" in t:
            return "pp_rf"
        return "pp"

    if "федеральный закон" in t:
        return "fed_law"

    if "закон" in t and ("г. москвы" in t or "города москвы" in t or "москвы" in t):
        return "law_moscow"

    if "закон рф" in t:
        return "law_rf"

    if "приказ" in t:
        return "order"

    if "распоряжение" in t:
        return "decree"

    return None


def _normalize_number(raw_num: str, tail: str) -> str:
    num = _norm_dashes(_norm_spaces(raw_num)).replace(" ", "")
    num = num.upper()

    # Handle cases like "N111 ПП" where suffix is separated.
    if re.fullmatch(r"[0-9]+(?:-[0-9]+)?", num):
        m = re.match(r"^\s*[-–—]?\s*([A-Za-zА-Яа-я]{1,8})\b", tail or "")
        if m:
            suf = m.group(1).upper()
            num = f"{num}-{suf}"

    return num


def extract_doc_keys(text: str) -> list[DocKey]:
    """Extract one or more document keys from free-form source strings."""
    s = _norm_dashes(text)
    if not s:
        return []

    doc_type = _detect_doc_type(s)
    if not doc_type:
        return []

    keys: list[DocKey] = []

    # Pattern 1: "от <date> ... № <num>"
    for m in re.finditer(
        r"(?i)\bот\s*(\d{2}\.\d{2}\.\d{4})\b(.{0,80}?)(?:№|\bN\b)\s*([0-9]+(?:[-–—][0-9]+)?(?:[-–—][A-Za-zА-Яа-я0-9]+)?)",
        s,
    ):
        date = m.group(1)
        raw_num = m.group(3)
        tail = s[m.end(3) : m.end(3) + 12]
        keys.append(
            DocKey(
                doc_type=doc_type, date=date, number=_normalize_number(raw_num, tail)
            )
        )

    # Pattern 2: "№ <num> ... от <date>" (some sources use this order)
    for m in re.finditer(
        r"(?i)(?:№|\bN\b)\s*([0-9]+(?:[-–—][0-9]+)?(?:[-–—][A-Za-zА-Яа-я0-9]+)?)\b(.{0,80}?)\bот\s*(\d{2}\.\d{2}\.\d{4})\b",
        s,
    ):
        raw_num = m.group(1)
        date = m.group(3)
        tail = s[m.end(1) : m.end(1) + 12]
        keys.append(
            DocKey(
                doc_type=doc_type, date=date, number=_normalize_number(raw_num, tail)
            )
        )

    # Fallback: best-effort if we only have one date and one number.
    if not keys:
        m_date = _OT_DATE_RE.search(s) or _DATE_RE.search(s)
        m_num = _NUM_RE.search(s)
        if m_date and m_num:
            date = m_date.group(1)
            raw_num = m_num.group(1)
            tail = s[m_num.end(1) : m_num.end(1) + 12]
            keys.append(
                DocKey(
                    doc_type=doc_type,
                    date=date,
                    number=_normalize_number(raw_num, tail),
                )
            )

    # De-dup while keeping order
    seen: set[DocKey] = set()
    uniq: list[DocKey] = []
    for k in keys:
        if k in seen:
            continue
        seen.add(k)
        uniq.append(k)
    return uniq


def build_available_doc_map(raw_dir: Path) -> dict[DocKey, str]:
    mapping: dict[DocKey, str] = {}
    for p in sorted(raw_dir.glob("*.docx")):
        keys = extract_doc_keys(p.stem)
        if not keys:
            continue
        # Usually exactly one key per filename.
        mapping.setdefault(keys[0], p.name)
    return mapping


def _safe_str(v) -> str:
    if v is None:
        return ""
    return str(v)


def _format_id(v) -> str:
    if v is None:
        return ""
    # openpyxl часто отдаёт числа как float
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    s = str(v).strip()
    if re.fullmatch(r"\d+\.0", s):
        return s[:-2]
    return s


def write_xlsx(path: Path, header: list[str], rows: Iterable[list[str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    ws.append(header)
    for r in rows:
        ws.append(r)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser(
        description=(
            "Prepare gold dataset from Excel and keep only questions whose source documents exist in data/raw_docx."
        )
    )
    ap.add_argument(
        "--xlsx",
        default="Вопросы для бота.xlsx",
        help="Path to input Excel (default: %(default)s)",
    )
    ap.add_argument(
        "--sheet",
        default="100вопросов_бота",
        help="Sheet name (default: %(default)s)",
    )
    ap.add_argument(
        "--raw-dir",
        default="data/raw_docx",
        help="Directory with raw .docx files (default: %(default)s)",
    )
    ap.add_argument(
        "--out-dir",
        default="data/gold",
        help="Output directory (default: %(default)s)",
    )
    args = ap.parse_args()

    xlsx_path = Path(args.xlsx)
    raw_dir = Path(args.raw_dir)
    out_dir = Path(args.out_dir)

    if not xlsx_path.exists():
        raise FileNotFoundError(f"Input Excel not found: {xlsx_path}")
    if not raw_dir.exists():
        raise FileNotFoundError(f"Raw docs directory not found: {raw_dir}")

    available = build_available_doc_map(raw_dir)

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    if args.sheet not in wb.sheetnames:
        raise KeyError(f"Sheet '{args.sheet}' not found. Available: {wb.sheetnames}")
    ws = wb[args.sheet]

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    col = {
        str(name).strip(): i for i, name in enumerate(header_row) if name is not None
    }
    for required in ("Вопрос", "Ответ", "Источник"):
        if required not in col:
            raise KeyError(
                f"Required column '{required}' not found in sheet '{args.sheet}'"
            )

    idx_id = col.get("№")
    idx_q = col["Вопрос"]
    idx_a = col["Ответ"]
    idx_src = col["Источник"]

    in_rows: list[list[str]] = []
    out_rows: list[list[str]] = []

    total = 0
    kept = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        total += 1

        q = _safe_str(row[idx_q]).strip()
        gold_answer = _safe_str(row[idx_a]).strip()
        src = _safe_str(row[idx_src]).strip()
        rid = _format_id(row[idx_id]) if idx_id is not None else ""

        if not q:
            continue

        if not src:
            out_rows.append([rid, q, gold_answer, "", src, "", "empty_source"])
            continue

        keys = extract_doc_keys(src)
        matched_name = ""
        for k in keys:
            if k in available:
                matched_name = available[k]
                break

        if matched_name:
            kept += 1
            in_rows.append([rid, q, gold_answer, "", src, matched_name])
        else:
            out_rows.append([rid, q, gold_answer, "", src, "", "source_not_in_raw_dir"])

    in_path = out_dir / "gold_in_scope.xlsx"
    out_path = out_dir / "gold_out_of_scope.xlsx"

    write_xlsx(
        in_path,
        header=[
            "id",
            "question",
            "gold_answer",
            "rag_answer",
            "gold_source",
            "matched_doc",
        ],
        rows=in_rows,
    )
    write_xlsx(
        out_path,
        header=[
            "id",
            "question",
            "gold_answer",
            "rag_answer",
            "gold_source",
            "matched_doc",
            "reason",
        ],
        rows=out_rows,
    )

    print(f"Input: {xlsx_path} sheet={args.sheet}")
    print(f"Raw docs: {raw_dir} (doc_keys={len(available)})")
    print(f"Rows processed: {total}")
    print(f"Kept (in-scope): {kept} -> {in_path}")
    print(f"Filtered (out-of-scope): {len(out_rows)} -> {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
